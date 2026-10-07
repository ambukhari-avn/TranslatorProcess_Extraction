"""Word-in / Word-out: units, formatting markers, write-back, headers, tables, fields."""
import os
import tempfile
import unittest
import zipfile

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

import docx_io
from masking import mask


def _sample(path):
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Note: ").bold = True
    p.add_run("the pump shall be inspected.")
    doc.add_paragraph("Plain paragraph with no formatting change.")
    doc.add_paragraph("12345")                                    # no letters: not a unit
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Wellhead"
    table.rows[0].cells[1].text = "Pressure"
    tr = table.rows[0]._tr
    height = tr.get_or_add_trPr().makeelement(docx_io.q("trHeight"), {docx_io.q("val"): "300", docx_io.q("hRule"): "exact"})
    tr.get_or_add_trPr().append(height)
    doc.sections[0].header.paragraphs[0].text = "Client procedure"
    tabbed = doc.add_paragraph("Name")
    tabbed.add_run("\tvalue")
    brk = doc.add_paragraph("first line")
    brk.add_run().add_break(WD_BREAK.LINE)
    brk.add_run("second line")
    doc.save(path)


class DocxTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.src = os.path.join(self.dir, "in.docx")
        _sample(self.src)

    def test_units_cover_body_table_and_header(self):
        texts = [s.text for s in docx_io.extract_docx(self.src)]
        self.assertIn("Plain paragraph with no formatting change.", texts)
        self.assertIn("Wellhead", texts)
        self.assertIn("Client procedure", texts)
        self.assertNotIn("12345", texts)
        kinds = {s.text: s.kind for s in docx_io.extract_docx(self.src)}
        self.assertEqual(kinds["Wellhead"], "table_cell")

    def test_a_formatting_change_becomes_a_marker_that_masking_protects(self):
        texts = [s.text for s in docx_io.extract_docx(self.src)]
        marked = next(t for t in texts if t.startswith("⟦1⟧Note"))
        self.assertEqual(marked, "⟦1⟧Note: ⟦2⟧the pump shall be inspected.")
        masked = mask(marked)
        self.assertNotIn("⟦", masked.masked_text)
        self.assertEqual(len(masked.mapping), 2)
        self.assertTrue(any("⟦t⟧" in t for t in texts) and any("⟦b⟧" in t for t in texts))

    def test_writing_the_same_text_back_changes_nothing_visible(self):
        out = os.path.join(self.dir, "same.docx")
        docx_io.translate_docx(self.src, out, lambda text, loc: (text, False))
        self.assertEqual([s.text for s in docx_io.extract_docx(self.src)], [s.text for s in docx_io.extract_docx(out)])

    def test_translation_keeps_the_bold_run_and_marks_flagged_text(self):
        out = os.path.join(self.dir, "ru.docx")
        table = {"⟦1⟧Note: ⟦2⟧the pump shall be inspected.": "⟦1⟧Примечание: ⟦2⟧насос должен проверяться.",
                 "Wellhead": "Устье скважины"}
        docx_io.translate_docx(self.src, out, lambda text, loc: (table.get(text, text), text == "Wellhead"))
        doc = Document(out)
        first = doc.paragraphs[0]
        self.assertEqual(first.text, "Примечание: насос должен проверяться.")
        self.assertTrue(first.runs[0].bold)
        self.assertFalse(first.runs[1].bold)
        cell_runs = doc.tables[0].rows[0].cells[0].paragraphs[0].runs
        self.assertEqual(cell_runs[0].text, "Устье скважины")
        with zipfile.ZipFile(out) as z:
            xml = z.read("word/document.xml").decode("utf8")
        self.assertIn('w:highlight w:val="yellow"', xml)
        self.assertIn('w:lang w:val="ru-RU"', xml)

    def test_tabs_breaks_and_fixed_row_heights_are_handled(self):
        out = os.path.join(self.dir, "ru2.docx")
        docx_io.translate_docx(self.src, out, lambda text, loc: (text.replace("Name", "Имя").replace("first", "первая"), False))
        with zipfile.ZipFile(out) as z:
            xml = z.read("word/document.xml").decode("utf8")
        self.assertIn("<w:tab/>", xml)
        self.assertIn("<w:br/>", xml)
        self.assertNotIn('w:hRule="exact"', xml)
        self.assertIn("Имя", xml)

    def test_other_parts_are_copied_untouched_and_a_dropped_marker_cannot_lose_text(self):
        out = os.path.join(self.dir, "ru3.docx")
        docx_io.translate_docx(self.src, out, lambda text, loc: ("Только текст" if "Note" in text else text, False))
        self.assertEqual(Document(out).paragraphs[0].text, "Только текст")
        with zipfile.ZipFile(self.src) as a, zipfile.ZipFile(out) as b:
            self.assertEqual(a.read("word/styles.xml"), b.read("word/styles.xml"))

    def test_a_footer_table_in_a_fixed_text_box_is_moved_into_the_footer(self):
        from docx.oxml import parse_xml
        footer = Document(self.src).sections[0].footer
        W = docx_io.W
        xml = (f'<w:r xmlns:w="{W}" xmlns:mc="{docx_io.MC}" xmlns:wp="{docx_io.WP}" xmlns:wps="{docx_io.WPS}" xmlns:a="{docx_io.DML}">'
               '<mc:AlternateContent><mc:Choice Requires="wps"><w:drawing><wp:anchor><wp:positionH relativeFrom="page"><wp:posOffset>635000</wp:posOffset></wp:positionH>'
               '<wp:positionV relativeFrom="page"><wp:posOffset>9000000</wp:posOffset></wp:positionV><wp:extent cx="5000000" cy="400000"/><wp:wrapNone/>'
               '<a:graphic><a:graphicData><wps:wsp><wps:txbx><w:txbxContent><w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="4000"/></w:tblGrid>'
               '<w:tr><w:tc><w:p><w:r><w:t>Copy Status:</w:t></w:r></w:p></w:tc></w:tr></w:tbl><w:p/></w:txbxContent></wps:txbx>'
               '<wps:bodyPr><a:noAutofit/></wps:bodyPr></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
               '<mc:Fallback><w:pict/></mc:Fallback></mc:AlternateContent></w:r>')
        footer.is_linked_to_previous = False
        footer.paragraphs[0]._p.append(parse_xml(xml))
        withbox = os.path.join(self.dir, "box.docx")
        Document(self.src).save(withbox)
        doc = Document(self.src)
        doc.sections[0].footer.is_linked_to_previous = False
        doc.sections[0].footer.paragraphs[0]._p.append(parse_xml(xml))
        doc.save(withbox)
        out = os.path.join(self.dir, "box_ru.docx")
        stats = docx_io.translate_docx(withbox, out, lambda text, loc: ("Статус копии:" if text == "Copy Status:" else text, False))
        self.assertEqual(stats["unwrapped"], 1)
        with zipfile.ZipFile(out) as z:
            footer_xml = z.read([n for n in z.namelist() if n.startswith("word/footer")][0]).decode("utf8")
            doc_xml = z.read("word/document.xml").decode("utf8")
        self.assertNotIn("txbxContent", footer_xml)
        self.assertIn("Статус копии:", footer_xml)
        self.assertIn("<w:tbl>", footer_xml)
        self.assertRegex(doc_xml, r'w:footer="\d+"')


if __name__ == "__main__":
    unittest.main()

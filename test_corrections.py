import unittest
import tempfile
from pathlib import Path
from contextlib import redirect_stdout
from io import StringIO
from unittest.mock import patch

from config import PipelineConfig
from extraction import Segment, _ocr_fallback
from knowledge_base import KnowledgeBase, TMEntry
from llm_client import build_prompt
from pipeline import process_segment, _gloss_sig, run_pipeline

from masking import mask, unmask
from qa_checks import run_checks


class PreservationTests(unittest.TestCase):
    def test_label_hyphens_are_not_negative_numbers(self):
        issues = run_checks("Refer Annexure-1", "\u0421\u043c. \u041f\u0440\u0438\u043b\u043e\u0436\u0435\u043d\u0438\u0435-1", [], [])
        self.assertNotIn("numeric_mismatch", [issue.rule for issue in issues])

    def test_chinese_next_to_translated_words_is_unchanged(self):
        source = "Colour\u989c\u8272"
        target = "\u0426\u0432\u0435\u0442\u989c\u8272"
        self.assertFalse(run_checks(source, target, [], [], list(mask(source).mapping.values())))

    def test_acronyms_do_not_lower_translated_word_ratio(self):
        source = "MTCs type 3.1, IMIR"
        target = "MTCs \u0442\u0438\u043f 3.1, IMIR"
        self.assertFalse(run_checks(source, target, [], [], list(mask(source).mapping.values())))

    def test_complete_identifiers_are_masked(self):
        for text in ("M-XX-XXX-001", "MI-TS", "RZK-NSS-LT- MI-U10-\n001", "RZK-NSS-PT-MI-U10-001/2", "EN 10204/3.1", "HART", "SST", "LL", "HH"):
            with self.subTest(text=text):
                result = mask(text)
                self.assertEqual(result.masked_text, "[[P0]]")
                self.assertEqual(unmask(result.masked_text, result.mapping), (text, []))

    def test_chinese_is_hidden_but_english_is_translatable(self):
        source = "\u76f8\u5bf9\u6e7f\u5ea6\nRelative humidity"
        result = mask(source)
        self.assertNotIn("\u76f8\u5bf9\u6e7f\u5ea6", result.masked_text)
        self.assertIn("Relative humidity", result.masked_text)
        translated, missing = unmask(result.masked_text.replace("Relative humidity", "\u041e\u0442\u043d\u043e\u0441\u0438\u0442\u0435\u043b\u044c\u043d\u0430\u044f \u0432\u043b\u0430\u0436\u043d\u043e\u0441\u0442\u044c"), result.mapping)
        self.assertEqual(translated.splitlines()[0], source.splitlines()[0])
        self.assertEqual(missing, [])
        self.assertFalse(run_checks(source, translated, [], [], list(result.mapping.values())))

    def test_missing_or_invented_chinese_is_critical(self):
        for source, target in (("\u76f8\u5bf9\u6e7f\u5ea6", "humidity"), ("humidity", "\u76f8\u5bf9\u6e7f\u5ea6")):
            issues = run_checks(source, target, [], [])
            self.assertIn("chinese_mismatch", [issue.rule for issue in issues])

    def test_chinese_adjacent_numbers_are_not_false_positives(self):
        issues = run_checks("\u7b2c11\u8282(11.1)", "Section 11 (11.1)", [], [])
        self.assertNotIn("numeric_mismatch", [issue.rule for issue in issues])

    def test_dropped_minus_remains_critical(self):
        issues = run_checks("-30\u00b0C", "30\u00b0C", [], [])
        self.assertIn("numeric_mismatch", [issue.rule for issue in issues])

    def test_changed_identifiers_are_critical(self):
        for source, target in (("MI-TS", "\u041c\u0418-TS"), ("M-XX-XXX-001", "\u041c-XX-XXX-001"), ("EN 10204/3.1", "RU 10204/3.1")):
            issues = run_checks(source, target, [], [], list(mask(source).mapping.values()))
            self.assertIn("protected_token_mismatch", [issue.rule for issue in issues])

    def test_hyphenated_english_is_not_a_tag(self):
        self.assertIn("HIGH-HIGH", mask("HIGH-HIGH PRESSURE").masked_text)

    def test_runaway_output_is_critical(self):
        issues = run_checks("Testing fluid", "\u0432\u043e\u0434\u0430 " * 200, [], [])
        self.assertIn("suspiciously_long_output", [issue.rule for issue in issues])


class PipelinePolicyTests(unittest.TestCase):
    def setUp(self):
        self.config = PipelineConfig()
        self.kb = KnowledgeBase("missing-tm.json", "missing-glossary.json")

    def test_chinese_only_needs_no_model(self):
        segment = Segment("\u7b2c11\u8282", 1, "paragraph", "page1/block0")
        with patch("llm_client.translate_segment", side_effect=AssertionError("Unexpected API call")):
            result = process_segment(segment, self.kb, self.config, [])
        self.assertEqual(result["translated_text"], segment.text)
        self.assertFalse(any(issue.severity == "critical" for issue in result["qa_issues"]))

    def test_tm_cannot_bypass_identifier_qa(self):
        self.kb.tm = [TMEntry("Tag MI-TS", "Tag \u041c\u0418-TS")]
        segment = Segment("Tag MI-TS", 1, "paragraph", "page1/block0")
        result = process_segment(segment, self.kb, self.config, [])
        self.assertIn("protected_token_mismatch", [issue.rule for issue in result["qa_issues"]])

    def test_prompt_preserves_chinese_in_all_modes(self):
        for options in ({}, {"minimal": True}, {"repair_text": "draft"}):
            prompt = build_prompt("[[P0]]\nRelative humidity", self.config, [], [], **options)
            self.assertIn("Preserve Chinese", prompt)
            self.assertNotIn("no Chinese characters at all", prompt)
            self.assertIn("English", prompt)

    def test_cache_signature_changes_with_tm_and_settings(self):
        initial = _gloss_sig(self.kb, "Testing fluid", self.config)
        self.kb.tm.append(TMEntry("Testing fluid", "\u0416\u0438\u0434\u043a\u043e\u0441\u0442\u044c"))
        changed_tm = _gloss_sig(self.kb, "Testing fluid", self.config)
        self.assertNotEqual(initial, changed_tm)
        self.config.temperature = 0.4
        self.assertNotEqual(changed_tm, _gloss_sig(self.kb, "Testing fluid", self.config))


class TerminologyTests(unittest.TestCase):
    def setUp(self):
        self.kb = KnowledgeBase("data/translation_memory.json", "data/glossary.json")

    def test_ambiguous_terms_do_not_apply_inside_other_labels(self):
        for text, unwanted in (("Kick off Meeting", "off"), ("Product Name:", "name:")):
            terms = self.kb.find_glossary_terms(text)
            self.assertNotIn(unwanted, [term.en_term.lower() for term in terms])

    def test_corrected_technical_glossary(self):
        expected = {
            "male": "\u043d\u0430\u0440\u0443\u0436\u043d\u0430\u044f \u0440\u0435\u0437\u044c\u0431\u0430",
            "of span": "\u043e\u0442 \u0434\u0438\u0430\u043f\u0430\u0437\u043e\u043d\u0430 \u0438\u0437\u043c\u0435\u0440\u0435\u043d\u0438\u044f",
            "Make": "\u041f\u0440\u043e\u0438\u0437\u0432\u043e\u0434\u0438\u0442\u0435\u043b\u044c",
            "Stem": "\u0428\u0442\u043e\u043a",
            "Face to face dimension": "\u0421\u0442\u0440\u043e\u0438\u0442\u0435\u043b\u044c\u043d\u0430\u044f \u0434\u043b\u0438\u043d\u0430",
            "HART": "HART",
        }
        for source, target in expected.items():
            with self.subTest(source=source):
                term = self.kb.glossary_exact(source)
                self.assertIsNotNone(term)
                self.assertEqual(term.target_term, target)

    def test_duplicate_placeholders_are_rejected(self):
        result = mask("MI-TS")
        _, invalid = unmask("[[P0]] [[P0]]", result.mapping)
        self.assertIn("[[P0]]", invalid)

    def test_legitimate_painting_names_and_acronyms(self):
        source = "JOTUN Testex MSDS WFT DFT SSPC DEPAMU ESI RZH AVANCEAON ppm OFT"
        issues = run_checks(source, source, [], [], list(mask(source).mapping.values()))
        self.assertNotIn("latin_words_in_output", [issue.rule for issue in issues])

    def test_full_company_name_is_preserved_not_globally_whitelisted(self):
        source = "VHV (CHINA VIHUNG VALVE CO, LTD."
        self.assertEqual(mask(source).masked_text, "[[P0]]")
        self.assertFalse(run_checks(source, source, [], [], list(mask(source).mapping.values())))

    def test_safety_data_sheet_has_specific_glossary(self):
        terms = self.kb.find_glossary_terms("Material Safety Data Sheet (MSDS)")
        self.assertIn("Material Safety Data Sheet", [term.en_term for term in terms])
        self.assertNotIn("Data sheet", [term.en_term for term in terms])

    def test_uppercase_roles_still_need_translation(self):
        source = "CONTRACTOR MANUFACTURER"
        issues = run_checks(source, source, [], [], list(mask(source).mapping.values()))
        self.assertIn("latin_words_in_output", [issue.rule for issue in issues])

    def test_chinese_span_cannot_gain_spaces(self):
        source = "\u7528\u6e7f\u5ea6\u8ba1\u6d4b\u91cf"
        target = "\u7528\u6e7f\u5ea6\u8ba1\u6d4b \u91cf"
        issues = run_checks(source, target, [], [], list(mask(source).mapping.values()))
        self.assertIn("protected_token_mismatch", [issue.rule for issue in issues])


class OutputTests(unittest.TestCase):
    def test_ocr_failure_reports_the_actual_cause(self):
        segment = _ocr_fallback(None, 6)[0]
        self.assertTrue(segment.text.startswith("[OCR FAILED:"))
        self.assertNotIn("NOT YET WIRED", segment.text)

    def test_bilingual_pdf_outputs_and_cache_without_network(self):
        import pymupdf
        from docx import Document

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            pdf_path = root / "bilingual.pdf"
            with pymupdf.open() as pdf:
                page = pdf.new_page()
                page.insert_text((72, 72), "\u76f8\u5bf9\u6e7f\u5ea6", fontname="china-s")
                page.insert_text((72, 100), "Relative humidity")
                pdf.save(pdf_path)
            config = PipelineConfig(workers=1, verify_translations=False, tm_path=str(root / "tm.json"), glossary_path=str(root / "glossary.json"))

            def respond(*args, **kwargs):
                content = kwargs["json"]["messages"][0]["content"]
                source = content.split("Text to translate:\n", 1)[1].split("\n\nRespond", 1)[0]
                text = source.replace("Relative humidity", "\u041e\u0442\u043d\u043e\u0441\u0438\u0442\u0435\u043b\u044c\u043d\u0430\u044f \u0432\u043b\u0430\u0436\u043d\u043e\u0441\u0442\u044c")
                from unittest.mock import Mock
                return Mock(status_code=200, ok=True, json=lambda: {"choices": [{"message": {"content": text}}]})

            config.openrouter_api_key = "offline-test-key"
            with patch("requests.post", side_effect=respond), redirect_stdout(StringIO()):
                results = run_pipeline(str(pdf_path), config, str(root / "output"))
            self.assertTrue(any("\u76f8\u5bf9\u6e7f\u5ea6" in item["translated_text"] for item in results))
            self.assertTrue(any("\u0432\u043b\u0430\u0436\u043d\u043e\u0441\u0442\u044c" in item["translated_text"] for item in results))
            review = Document(root / "output" / "bilingual_ru_review.docx")
            self.assertGreater(len(review.tables[0].rows), 1)
            from pipeline import _cjk_in_docx
            self.assertIn("\u76f8\u5bf9\u6e7f\u5ea6", _cjk_in_docx(root / "output" / "bilingual_ru_replica.docx"))
            with patch("requests.post", side_effect=AssertionError("Cache should avoid network")), redirect_stdout(StringIO()):
                run_pipeline(str(pdf_path), config, str(root / "output"))

    def test_failed_release_saves_review_before_raising(self):
        import pymupdf

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            pdf_path = root / "failed.pdf"
            with pymupdf.open() as pdf:
                pdf.new_page().insert_text((72, 72), "Missing translation")
                pdf.save(pdf_path)
            config = PipelineConfig(workers=1, verify_translations=False, tm_path=str(root / "tm.json"), glossary_path=str(root / "glossary.json"))
            with patch("llm_client.translate_segment", side_effect=RuntimeError("Offline failure")), redirect_stdout(StringIO()):
                with self.assertRaisesRegex(RuntimeError, "REVIEW REQUIRED"):
                    run_pipeline(str(pdf_path), config, str(root / "output"))
            self.assertTrue((root / "output" / "failed_ru_review.docx").exists())


if __name__ == "__main__":
    unittest.main()
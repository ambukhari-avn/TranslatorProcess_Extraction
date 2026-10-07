"""Writes the review Word file: one row per segment, shaded when it came from OCR or failed a QA check."""

import os

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


HIGHLIGHT_OCR = "FFF2CC"       # pale yellow: from OCR
HIGHLIGHT_ISSUE = "F8CBAD"     # pale red: moderate QA issue
HIGHLIGHT_CRITICAL = "FF6B6B"  # strong red: critical QA issue


def _shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def write_review_document(segments_with_results: list, output_path: str,
                           doc_title: str = "Translation Review", verification_lines=None):
    """Review table (English, translation, notes), shaded by QA severity. Each item is a dict with
    "segment", "translated_text", "match_type" and "qa_issues"."""
    doc = Document()
    doc.add_heading(doc_title, level=1)

    # summary first: only what needs a human
    flagged = [(it["segment"].location, qa) for it in segments_with_results for qa in it.get("qa_issues", [])
               if qa.severity in ("critical", "moderate")]
    flagged.sort(key=lambda x: (x[1].severity != "critical", x[0]))
    n_crit = sum(1 for _, q in flagged if q.severity == "critical")
    doc.add_paragraph("REVIEW REQUIRED - not approved for release." if flagged
                      else "Automated QA passed; engineering review is still required.")
    doc.add_paragraph(f"{n_crit} critical and {len(flagged) - n_crit} moderate issue(s). "
                      f"Critical first; every row is also marked in the table below.")
    for line in verification_lines or []:
        doc.add_paragraph(line)
    for loc, qa in flagged[:40]:
        doc.add_paragraph(f"[{qa.severity.upper()}] {loc} - {qa.rule}: {qa.message}", style="List Bullet")
    if len(flagged) > 40:
        doc.add_paragraph(f"... and {len(flagged) - 40} more in the table.")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Location"
    hdr[1].text = "English (source)"
    hdr[2].text = "Translation"
    hdr[3].text = "Notes"

    for item in segments_with_results:
        seg = item["segment"]
        row = table.add_row().cells
        row[0].text = seg.location
        row[1].text = seg.text
        row[2].text = item["translated_text"]

        notes = []
        if seg.from_ocr:
            notes.append("from OCR — verify against source")
        if item["match_type"] == "exact":
            notes.append("reused from approved TM (not AI-translated)")
        elif item["match_type"] == "fuzzy":
            notes.append("guided by similar approved example")

        if item["match_type"] == "glossary":
            notes.append("glossary term used directly (not AI-translated)")
        if item["match_type"] == "rule":
            notes.append("fixed pattern (not AI-translated)")
        confidence = (item.get("verification") or {}).get("confidence")
        if confidence is not None:
            notes.append(f"verifier confidence {confidence}/100")
        max_severity = None
        for qa in item.get("qa_issues", []):
            if qa.rule == "kept_as_is":
                continue
            notes.append(f"[{qa.severity.upper()}] {qa.rule}: {qa.message}")
            if qa.severity == "critical":
                max_severity = "critical"
            elif qa.severity == "moderate" and max_severity != "critical":
                max_severity = "moderate"

        row[3].text = "; ".join(notes) if notes else ""

        if max_severity == "critical":
            for c in row:
                _shade_cell(c, HIGHLIGHT_CRITICAL)
        elif max_severity == "moderate":
            for c in row:
                _shade_cell(c, HIGHLIGHT_ISSUE)
        elif seg.from_ocr:
            for c in row:
                _shade_cell(c, HIGHLIGHT_OCR)

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)